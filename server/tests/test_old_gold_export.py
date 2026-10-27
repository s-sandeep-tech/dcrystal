import ast
import json
import os
import tempfile
import unittest
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path
from unittest.mock import Mock, patch

from flask import Flask
from flask_jwt_extended import JWTManager, create_access_token
from openpyxl import load_workbook
from sqlalchemy.sql.elements import Cast
from sqlalchemy.sql.visitors import replacement_traverse

from app.extensions import db
from app.models import LocationWiseOldGoldSettlementTransferSnapshot as Snapshot
from app.dashboard.routes import location_wise_old_gold_settlement_transfer as report
from app.services.old_gold_export import generate_location_wise_old_gold_export


class OldGoldExportTests(unittest.TestCase):
    def setUp(self):
        self.app = Flask(__name__)
        self.app.config.update(TESTING=True, SECRET_KEY='test-secret-at-least-32-characters',
                               JWT_SECRET_KEY='test-jwt-secret-at-least-32-characters',
                               SQLALCHEMY_DATABASE_URI='sqlite:///:memory:')
        db.init_app(self.app)
        JWTManager(self.app)
        self.ctx = self.app.app_context()
        self.ctx.push()
        Snapshot.__table__.create(db.engine)
        self.app.add_url_rule('/api/export', view_func=report.queue_location_wise_old_gold_export, methods=['POST'])
        self.client = self.app.test_client()
        self.headers = {'Authorization': 'Bearer ' + create_access_token(identity='7')}
        self.as_of = date(2026, 10, 27)
        self.tmp = tempfile.TemporaryDirectory()
        self.paths = patch('app.utils.export_service.EXPORTS_DIR', self.tmp.name)
        self.paths.start()
        original = report.get_aggregation_columns
        # SQLite CAST(date AS DATE) is numeric, unlike PostgreSQL's DATE cast.
        self.date_cast = patch.object(report, 'get_aggregation_columns', side_effect=lambda as_of=None: [
            replacement_traverse(c, {}, lambda n: n.clause if isinstance(n, Cast) else None)
            for c in original(as_of)])
        self.date_cast.start()

    def tearDown(self):
        self.date_cast.stop()
        self.paths.stop()
        db.session.remove()
        self.ctx.pop()
        self.tmp.cleanup()

    def add(self, days=1, **changes):
        row = dict(id=Snapshot.query.count() + 1, office='RO A', locationname='SHOP A',
                   transdate=self.as_of - timedelta(days=days) if days is not None else None,
                   grwt=Decimal('1.125'), intransit_grwt=Decimal('2.250'), transfer_grwt=999,
                   division='Gold', groupname='OLD GOLD', purity='920', locationtype='SHOP')
        row.update(changes)
        db.session.add(Snapshot(**row))
        db.session.commit()

    def workbook(self, filters=None):
        name = generate_location_wise_old_gold_export(filters or {}, self.as_of)
        workbook = load_workbook(Path(self.tmp.name) / name)
        self.addCleanup(workbook.close)
        return workbook.active

    def test_headers_all_bucket_boundaries_and_totals(self):
        for age in (1, 2, 5, 6, 10, 11, 15, 16):
            self.add(days=age)
        sheet = self.workbook()
        self.assertEqual(sheet.max_column, 14)
        self.assertEqual([sheet.cell(1, c).value for c in (1, 2, 3, 5, 7, 9, 11, 13)],
                         ['RO NAME', 'SHOP NAME', '1 DAY', '2 - 5 DAYS', '6 - 10 DAYS', '11 - 15 DAYS', '> 15 DAYS', 'TOTAL'])
        self.assertEqual([sheet.cell(2, c).value for c in range(3, 15)], ['PENDING', 'INTRANSIT'] * 6)
        self.assertEqual([sheet.cell(3, c).value for c in range(3, 15)],
                         [1.125, 2.25, 2.25, 4.5, 2.25, 4.5, 2.25, 4.5, 1.125, 2.25, 9, 18])
        self.assertEqual(sheet['A4'].value, 'GRAND TOTAL')
        self.assertEqual(sheet['M4'].value, 9)
        self.assertEqual(sheet['C3'].data_type, 'n')
        self.assertEqual(sheet['C3'].number_format, '#,##0.000')
        self.assertEqual(sheet.freeze_panes, 'C3')

    def test_grouping_is_ro_and_shop_not_page_or_product(self):
        self.add()
        self.add(division='Silver')
        self.add(office='RO B')
        self.add(locationname='SHOP B')
        sheet = self.workbook({'page': '2', 'per_page': '1'})
        self.assertEqual(sheet.max_row, 6)
        self.assertEqual(sheet['C3'].value, 2.25)
        self.assertEqual(sheet['M6'].value, 4.5)

    def test_all_applied_filters_are_shared_with_report(self):
        self.add()
        self.add(office='RO B', locationname='SHOP B')
        self.add(division='Silver')
        self.add(days=5)
        sheet = self.workbook({'office': 'RO A', 'location': 'SHOP A', 'locationtype': 'SHOP',
            'division': 'Gold', 'group': 'OLD GOLD', 'purity': '920', 'search': 'SHOP A',
            'enable_date_filter': 'true', 'from_date': '2026-10-26', 'to_date': '2026-10-26'})
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual(sheet['M3'].value, 1.125)

    def test_null_today_future_dates_match_overall_report_totals(self):
        for age in (None, 0, -1):
            self.add(days=age)
        sheet = self.workbook()
        self.assertTrue(all(sheet.cell(3, c).value == 0 for c in range(3, 13)))
        self.assertEqual(sheet['M3'].value, 3.375)
        self.assertEqual(sheet['N3'].value, 6.75)

    def test_empty_and_formula_like_labels(self):
        sheet = self.workbook()
        self.assertEqual(sheet['A3'].value, 'GRAND TOTAL')
        self.assertEqual(sheet['M3'].value, 0)
        self.add(office='=1+1', locationname='+SUM(A1:A2)')
        sheet = self.workbook()
        self.assertEqual(sheet['A3'].data_type, 's')
        self.assertEqual(sheet['B3'].data_type, 's')

    def test_queue_requires_auth_and_permission(self):
        self.assertEqual(self.client.post('/api/export', json={}).status_code, 401)
        with patch('app.utils.decorators.get_user_permissions', return_value=set()):
            self.assertEqual(self.client.post('/api/export', json={}, headers=self.headers).status_code, 403)

    def test_queue_whitelists_filters_ignores_forged_identity_and_pagination(self):
        with patch('app.utils.decorators.get_user_permissions', return_value={'report.export'}), \
             patch.object(report, 'redis_client') as redis:
            response = self.client.post('/api/export', headers=self.headers, json={
                'user_id': 'someone-else', 'as_of': '2000-01-01', 'socket_id': 'socket-1',
                'filters': {'office': 'RO A', 'page': '2', 'per_page': '1'}})
        self.assertEqual(response.status_code, 200)
        queue, raw = redis.rpush.call_args.args
        payload = json.loads(raw)
        self.assertEqual(queue, 'export_queue')
        self.assertEqual(payload['type'], 'export_location_wise_old_gold')
        self.assertEqual(payload['user_id'], '7')
        self.assertEqual(payload['filters'], {'office': 'RO A'})
        self.assertNotEqual(payload['as_of'], '2000-01-01')

    def test_queue_rejects_bad_dates_and_filter_types(self):
        with patch('app.utils.decorators.get_user_permissions', return_value={'report.export'}), \
             patch.object(report, 'redis_client') as redis:
            for data in ({'filters': []}, {'filters': {'office': ['RO A']}},
                         {'filters': {'enable_date_filter': 'true', 'from_date': '2026-10-28', 'to_date': '2026-10-27'}}):
                self.assertEqual(self.client.post('/api/export', json=data, headers=self.headers).status_code, 400)
            redis.rpush.assert_not_called()

    def test_queue_failure_is_reported(self):
        with patch('app.utils.decorators.get_user_permissions', return_value={'report.export'}), \
             patch.object(report, 'redis_client') as redis:
            redis.rpush.side_effect = RuntimeError('redis unavailable')
            self.assertEqual(self.client.post('/api/export', json={}, headers=self.headers).status_code, 503)

    def test_worker_success_and_failure_notifications(self):
        consumer = Path(os.environ.get('SYNC_CONSUMER_SOURCE',
            Path(__file__).resolve().parents[2] / 'sync-consumer/main.py'))
        if not consumer.exists():
            self.skipTest('Set SYNC_CONSUMER_SOURCE when testing outside the repository root')
        function = next(n for n in ast.parse(consumer.read_text()).body
                        if isinstance(n, ast.FunctionDef) and n.name == '_handle_export_location_wise_old_gold')
        namespace = {'logger': Mock(), 'redis_client': Mock(), 'json': json}
        exec(compile(ast.Module(body=[function], type_ignores=[]), str(consumer), 'exec'), namespace)
        for fail in (False, True):
            with patch('app.services.old_gold_export.generate_location_wise_old_gold_export') as generate, \
                 patch('app.extensions.db') as database, patch('app.extensions.socketio') as socket:
                generate.return_value = 'old_gold_test.xlsx'
                if fail:
                    generate.side_effect = ValueError('generation failed')
                namespace['_handle_export_location_wise_old_gold']({'filters': {}, 'user_id': '7', 'socket_id': 's'})
                notification = database.session.add.call_args.args[0]
                self.assertEqual(notification.notification_type, 'error' if fail else 'success')
                self.assertEqual(notification.user_id, '7')
                self.assertEqual(notification.action_url, None if fail else '/exports/download/old_gold_test.xlsx')
                self.assertEqual(socket.emit.call_args.args[1]['user_id'], '7')


if __name__ == '__main__':
    unittest.main()
