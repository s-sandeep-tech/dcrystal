import os
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.services.location_provision_stock_analysis_data import (
    LocationProvisionStockAnalysisData,
)


class LocationProvisionStockAnalysisExport:
    """Excel export owned exclusively by Location Provision & Stock Analysis."""

    EXPORTS_DIR = '/app/uploads/exports'

    @classmethod
    def generate(cls, filters):
        os.makedirs(cls.EXPORTS_DIR, exist_ok=True)
        params = {
            key: filters.get(key) or None
            for key in (
                'location', 'state', 'purity', 'classification', 'make',
                'collection', 'section', 'prov_type', 'provision_mode',
                'branch_type', 'branch_status', 'business_head',
                'bh_emp_code', 'authorized_branch_ids',
            )
        }
        rows = LocationProvisionStockAnalysisData.fetch_summary_rows(params)

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Provision Stock Analysis'
        headers = (
            'Report Section', 'Detail', 'Provision Pcs', 'Provision Gross Wt',
            'In Shop Wt', 'Transit Wt', 'Short Pcs', 'Excess Pcs', 'Short %', 'Short Wt',
            'Excess Wt', 'Net Short / Excess', 'Ordered Wt',
        )
        sheet.append(headers)

        header_fill = PatternFill('solid', fgColor='E8EEF7')
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='24324A')
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

        for row in rows:
            sheet.append((
                row.get('report_section'),
                row.get('report_label'),
                row.get('prov_pcs'),
                row.get('prov_gr_wt'),
                row.get('in_shop_wt'),
                row.get('in_transit_wt'),
                row.get('short_pcs'),
                row.get('excess_pcs'),
                row.get('short_percent'),
                row.get('short_wt'),
                row.get('excess_wt'),
                (row.get('excess_wt') or 0) - (row.get('short_wt') or 0),
                row.get('ordered_wt'),
            ))

        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        widths = (22, 36, 15, 20, 16, 16, 14, 14, 12, 16, 16, 20, 16)
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + index)].width = width

        timestamp = datetime.now(ZoneInfo('Asia/Kolkata')).strftime('%Y%m%d_%H%M%S')
        filename = f'location_provision_stock_analysis_{timestamp}.xlsx'
        workbook.save(os.path.join(cls.EXPORTS_DIR, filename))
        return filename
