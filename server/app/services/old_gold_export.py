"""Shop-level old gold export using the report's filters and aging expressions."""
import os
from datetime import date
from decimal import Decimal
from uuid import uuid4
from sqlalchemy import func

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.extensions import db
from app.models import LocationWiseOldGoldSettlementTransferSnapshot as Snapshot
from app.utils import export_service


BUCKETS = [('1 DAY', 'today'), ('2 - 5 DAYS', '2_5'), ('6 - 10 DAYS', '6_10'),
           ('11 - 15 DAYS', '11_15'), ('> 15 DAYS', 'gt_15'), ('TOTAL', 'total')]


def generate_location_wise_old_gold_export(filters, as_of=None):
    from app.dashboard.routes.location_wise_old_gold_settlement_transfer import (
        build_filters, get_aggregation_columns, get_sort_order_clause,
    )

    as_of = date.fromisoformat(as_of) if isinstance(as_of, str) else as_of
    as_of = as_of or db.session.query(func.current_date()).scalar()
    fields = [field for _, bucket in BUCKETS for field in
              [('tot_grwt' if bucket == 'total' else 'grwt_' + bucket), 'intransit_' + bucket]]
    columns = {column.name: column for column in get_aggregation_columns(as_of)}
    groups = [Snapshot.office, Snapshot.locationname]
    query = build_filters(db.session.query(Snapshot), filters)
    rows = query.with_entities(*groups, *(columns[field] for field in fields)).group_by(*groups).order_by(
        *get_sort_order_clause(filters.get('sort_by'), filters.get('sort_order'), groups)
    ).yield_per(1000)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Old Gold Pending'
    sheet.sheet_view.showGridLines = False
    for col, label in [(1, 'RO NAME'), (2, 'SHOP NAME')]:
        sheet.cell(1, col, label)
        sheet.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
    for i, (label, _) in enumerate(BUCKETS):
        col = 3 + i * 2
        sheet.cell(1, col, label)
        sheet.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 1)
        sheet.cell(2, col, 'PENDING')
        sheet.cell(2, col + 1, 'INTRANSIT')
    for row in sheet.iter_rows(min_row=1, max_row=2, max_col=14):
        for cell in row:
            cell.fill = PatternFill('solid', fgColor='285E59')
            cell.font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
            cell.alignment = Alignment(horizontal='center', vertical='center')
    sheet.row_dimensions[1].height = 26
    sheet.row_dimensions[2].height = 24
    totals = [Decimal(0)] * 12
    for record in rows:
        values = [Decimal(str(getattr(record, field) or 0)) for field in fields]
        sheet.append([record.office or 'Unknown', record.locationname or 'Unknown', *values])
        # Source labels must remain text, even if they begin with Excel formula syntax.
        for col in (1, 2):
            sheet.cell(sheet.max_row, col).data_type = 's'
        totals = [total + value for total, value in zip(totals, values)]
    last_data_row = sheet.max_row
    sheet.append(['GRAND TOTAL', '', *totals])
    for row in sheet.iter_rows(min_row=3, max_col=14):
        for cell in row:
            cell.font = Font(name='Arial', size=11, bold=cell.row == sheet.max_row)
            cell.alignment = Alignment(horizontal='left' if cell.column < 3 else 'right', vertical='center')
            if cell.column >= 3:
                cell.number_format = '#,##0.000'
            if cell.row == sheet.max_row or cell.row % 2:
                cell.fill = PatternFill('solid', fgColor='EDF3F2')
            if cell.column in (4, 6, 8, 10, 12):
                cell.border = Border(right=Side(style='thin', color='ADBDBA'))
        sheet.row_dimensions[row[0].row].height = 23
    sheet.column_dimensions['A'].width = 38
    sheet.column_dimensions['B'].width = 38
    for col in range(3, 15):
        sheet.column_dimensions[get_column_letter(col)].width = 19
    sheet.freeze_panes = 'C3'
    if last_data_row >= 3:
        sheet.auto_filter.ref = f'A2:N{last_data_row}'
    sheet.print_title_rows = '1:2'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddHeader.center.text = f'Old Gold Pending / InTransit - As of {as_of.isoformat()}'
    export_service._ensure_exports_dir()
    filename = f'location_wise_old_gold_{uuid4().hex}.xlsx'
    workbook.save(os.path.join(export_service.EXPORTS_DIR, filename))
    return filename
