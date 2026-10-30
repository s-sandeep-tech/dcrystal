from io import BytesIO
from datetime import date
from flask import jsonify, render_template, request, send_file
from app.dashboard import dashboard_bp
from app.extensions import db
from app.utils.decorators import require_report_access, require_role, require_perm
from app.services.nip_stock_trend import DIMENSIONS, FILTER_FIELDS, scoped_query, trend_data
from app.models.nip_stock_trend import NIPStockTrendSnapshot as S

ROUTE = '/nip-stock-trend-analysis'


@dashboard_bp.route(ROUTE)
@require_report_access(ROUTE)
def nip_stock_trend_page():
    return render_template('nip_stock_trend.html', dimensions=DIMENSIONS, filter_fields=FILTER_FIELDS,
                           unread_count=0, sync_time='')


@dashboard_bp.route('/api/nip-stock-trend-analysis/options')
@require_report_access(ROUTE)
def nip_stock_trend_options():
    try:
        query = scoped_query()
        options = {}
        for name in FILTER_FIELDS:
            column = getattr(S, name)
            values = query.with_entities(column).filter(column.isnot(None))
            if name != 'weight':
                values = values.filter(column != '')
            options[name] = [format(row[0], 'f') if name == 'weight' else row[0]
                             for row in values.distinct().order_by(column)]
        from sqlalchemy import func
        first, last = query.with_entities(func.min(S.date), func.max(S.date)).one()
        return jsonify(options=options, first_date=first.isoformat() if first else None,
                       last_date=last.isoformat() if last else None)
    except Exception:
        db.session.rollback()
        return jsonify(error='NIP snapshot unavailable. Run NIP Stock Trend Analysis Sync in Settings.'), 503


@dashboard_bp.route('/api/nip-stock-trend-analysis/data')
@require_report_access(ROUTE)
def nip_stock_trend_data():
    try:
        return jsonify(trend_data(request.args))
    except ValueError as error:
        return jsonify(error=str(error)), 400
    except Exception:
        db.session.rollback()
        return jsonify(error='Unable to load NIP stock history. Check the sync status.'), 503


@dashboard_bp.route('/api/sync/nip-stock-trend-analysis', methods=['POST'])
@require_role(['ADMIN', 'DATA_SYNC_USER'])
def sync_nip_stock_trend():
    from app.utils.sync_manager import enqueue_sync_task
    from flask_jwt_extended import get_jwt_identity
    return jsonify(enqueue_sync_task('nip_stock_trend', get_jwt_identity()))


@dashboard_bp.route('/api/nip-stock-trend-analysis/export')
@require_report_access(ROUTE)
@require_perm('report.export')
def export_nip_stock_trend():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side
    try:
        data = trend_data(request.args)
    except ValueError as error:
        return jsonify(error=str(error)), 400
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'NIP Stock Trend'
    sheet.append(['Period Start', 'Balance Date', data['group_by'].replace('_', ' ').title(),
                  f"Total ({data['unit']})", f"NIP ({data['unit']})", f"Excess ({data['unit']})"])
    for row in data['table_rows']:
        sheet.append([date.fromisoformat(row['period']), date.fromisoformat(row['balance_date']),
                      row['group'], row['total'], row['nip'], row['excess']])
        sheet.cell(sheet.max_row, 3).data_type = 's'
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='E7E6E6')
        cell.font = Font(bold=True, color='333333')
    for row in sheet.iter_rows(min_row=2):
        row[0].number_format = row[1].number_format = 'yyyy-mm-dd'
        for cell in row[3:]:
            cell.number_format = '#,##0.000' if data['metric'] == 'gross_wt' else '#,##0'
        for cell in row:
            cell.border = Border(right=Side(style='thin', color='D0D0D0'))
    for column, width in [('A', 18), ('B', 18), ('C', 42), ('D', 22), ('E', 22), ('F', 22)]:
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = sheet.dimensions
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return send_file(buffer, as_attachment=True, download_name='nip_stock_trend.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
